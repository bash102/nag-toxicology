"""Build data/air_toxics_risk.geojson: U.S. EPA AirToxScreen 2020 modeled lifetime
cancer risk and respiratory hazard index for the census tracts of Brooklyn CD1.

Source: EPA Air Toxics Screening Assessment (AirToxScreen), 2020 assessment
(the newest year with risk results; EPA has released only emissions and
concentrations for 2021). The 2020 assessment reports results for **2020 census
blocks**, not tracts, so the block results are aggregated to 2020 tracts with
EPA's own method (population-weighted mean, TSD section 6.4.1) and joined on
GEOID to data/acs.geojson, whose tract geometry is reused.

Files (gaftp.epa.gov/rtrmodeling_public/AirToxScreen/2020/ and www.epa.gov):
  - Cancer/BySource/Region*_CancerRisk_by_block_srcgrp.xlsx (all 13, for the
    national and state percentiles)
  - Cancer/ByPollutant/Region2_CancerRisk_by_block_poll.xlsx
  - Exposure Concentrations/Region2_2020ATS_Exposure_Concentrations.xlsx plus
    the RfCs and target organs in 2020AirToxScreenPollutants.xlsx (supplemental
    data zip), for the respiratory hazard index, which EPA describes but does not
    publish as a file for 2020
  - national_cancerrisk_by_county_srcgrp.xlsx (EPA county averages, used as a check)

The workbooks are large (1.6 GB for the 13 source-group files). They are read with
python-calamine if it is installed (pip install python-calamine), which is ~7x
faster than the openpyxl fallback. Derived national tract totals are cached as
JSON so re-runs are quick.
"""
import json
import math
import os
import statistics
import zipfile
from collections import defaultdict

from common import CACHE_DIR, DATA_DIR, fetch, write_fc

YEAR = 2020
FTP = "https://gaftp.epa.gov/rtrmodeling_public/AirToxScreen/2020"
REGIONS = ["1", "2", "3", "4a", "4b", "5a", "5b", "6a", "6b", "7", "8", "9", "10"]
SRC_URL = FTP + "/Cancer/BySource/Region{}_CancerRisk_by_block_srcgrp.xlsx"
POLL_URL = FTP + "/Cancer/ByPollutant/Region2_CancerRisk_by_block_poll.xlsx"
EXP_URL = FTP + "/Exposure%20Concentrations/Region2_2020ATS_Exposure_Concentrations.xlsx"
SUPP_URL = ("https://www.epa.gov/system/files/other-files/2024-06/"
            "2020-airtoxscreen-supplemental-data-files.zip")
COUNTY_URL = ("https://www.epa.gov/system/files/documents/2025-01/"
              "national_cancerrisk_by_county_srcgrp.xlsx")

NYC_COUNTIES = {"36005", "36047", "36061", "36081", "36085"}
# 50 states + DC (state FIPS); Puerto Rico (72) and the Virgin Islands (78) are
# left out of the "national" percentile.
NON_STATES = {"72", "78"}

# EPA source-type columns -> broad source groups (TSD table 2-1 / mapping-tool totals)
GROUPS = [
    ("PT-", "Point sources"),
    ("OR-", "On-road mobile"),
    ("NR-", "Nonroad mobile"),
    ("NP-", "Nonpoint"),
    ("FIRE", "Fires"),
    ("BIOGENICS", "Biogenics"),
    ("SECONDARY", "Secondary formation"),
    ("BACKGROUND", "Background"),
]


# Readable labels for EPA's detailed source types (mapping-tool aliases, shortened)
SOURCE_TYPES = {
    "PT-StationaryPoint": "Stationary point sources",
    "OR-LightDuty-OffNetwork-Gas": "On-road light-duty gas (off-network)",
    "OR-LightDuty-OffNetwork-Diesel": "On-road light-duty diesel (off-network)",
    "OR-HeavyDuty-OffNetwork-Gas": "On-road heavy-duty gas (off-network)",
    "OR-HeavyDuty-OffNetwork-Diesel": "On-road heavy-duty diesel (off-network)",
    "OR-LightDuty-OnNetwork-Gas": "On-road light-duty gas (on-network)",
    "OR-LightDuty-OnNetwork-Diesel": "On-road light-duty diesel (on-network)",
    "OR-HeavyDuty-OnNetwork-Gas": "On-road heavy-duty gas (on-network)",
    "OR-HeavyDuty-OnNetwork-Diesel": "On-road heavy-duty diesel (on-network)",
    "OR-Refueling": "On-road refueling",
    "OR-HeavyDuty-Hoteling": "On-road heavy-duty hoteling",
    "NR-Recreational-inc-PleasureCraft": "Nonroad recreational (incl. pleasure craft)",
    "NR-Construction": "Nonroad construction equipment",
    "NR-CommercialLawnGarden": "Nonroad commercial lawn/garden",
    "NR-ResidentialLawnGarden": "Nonroad residential lawn/garden",
    "NR-Agriculture": "Nonroad agricultural equipment",
    "NR-CommercialEquipment": "Nonroad commercial equipment",
    "NR-AllOther": "Nonroad other equipment",
    "NR-CMV_C1C2_ports": "Commercial marine vessels C1/C2 (ports)",
    "NR-CMV_C3_ports": "Commercial marine vessels C3 (ports)",
    "NR-CMV_C1C2C3_underway": "Commercial marine vessels (underway)",
    "NR-Locomotives": "Locomotives",
    "NR-Point-Airports": "Airports",
    "NR- Point-Railyards": "Railyards",
    "NP-industrial": "Nonpoint industrial",
    "NP-CommercialCooking": "Commercial cooking",
    "NP-OilGas": "Nonpoint oil and gas",
    "NP-SolventsCoatings": "Solvents and coatings",
    "NP-StorageTransfer_BulkTerminals_GasStage1": "Storage/transfer: bulk terminals and gas stage 1",
    "NP-MiscellaneousNonindustrial": "Nonpoint misc. nonindustrial",
    "NP-FuelCombustion_not_RWC": "Fuel combustion (not residential wood)",
    "NP-ResidentialWoodCombustionRWC": "Residential wood combustion",
    "NP-WasteDisposal": "Waste disposal",
    "NP-AgricultureLivestockWaste": "Agricultural livestock waste",
    "NP-AgricultureLivestockSilage": "Agricultural livestock silage",
    "FIRE": "Fires",
    "BIOGENICS": "Biogenics",
    "SECONDARY": "Secondary formation",
    "BACKGROUND": "Background",
}


def source_label(col):
    return SOURCE_TYPES[col.replace("Cancer Risk (per million)", "").strip()]


def group_of(col):
    for prefix, name in GROUPS:
        if col.startswith(prefix):
            return name
    raise ValueError(f"unknown AirToxScreen source column {col!r}")


def cached_file(url):
    """Download `url` once into $NAG_CACHE under a readable name; return its path."""
    name = "airtoxscreen2020_" + url.rsplit("/", 1)[1].replace("%20", "_")
    path = os.path.join(CACHE_DIR, name)
    if not os.path.exists(path):
        print(f"  downloading {url}")
        raw = fetch(url, binary=True, cache=False, timeout=600)
        with open(path + ".part", "wb") as f:
            f.write(raw)
        os.replace(path + ".part", path)
    return path


def xlsx_rows(path):
    """Yield the rows of the first sheet of an .xlsx as tuples."""
    try:
        from python_calamine import CalamineWorkbook
    except ImportError:
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True)
        for r in wb.worksheets[0].iter_rows(values_only=True):
            yield tuple(r)
        wb.close()
        return
    ws = CalamineWorkbook.from_path(path).get_sheet_by_index(0)
    for r in ws.iter_rows():
        yield tuple(r)


def block_id(v):
    s = str(int(v)) if isinstance(v, float) else str(v).strip()
    return s.zfill(15)


def num(v):
    if v is None or v == "":
        return 0.0
    return float(v)


# --------------------------------------------------------------------------
# National tract totals (for percentiles and context medians)
# --------------------------------------------------------------------------

def national_tracts(cd1_ids):
    """Population-weighted tract means of total cancer risk for every US tract,
    from the 13 regional block-by-source-group workbooks.

    Block total = sum of EPA's source-type columns (the published "Total Cancer
    Risk" column is rounded to one significant figure). Also returns, for CD1
    tracts, population-weighted means of every source column."""
    cache = os.path.join(CACHE_DIR, "airtoxscreen2020_tract_totals.json")
    if os.path.exists(cache):
        with open(cache) as f:
            d = json.load(f)
        if d.get("regions") == REGIONS and set(d["cd1_detail"]) <= cd1_ids:
            return d
    tracts = {}       # geoid -> [pop, sum(pop*risk), n_blocks]
    detail = {}       # CD1 geoid -> {col: sum(pop*risk)}
    src_cols = None
    checks = {"blocks": 0, "max_abs_diff_vs_rounded_total": 0.0}
    for reg in REGIONS:
        path = cached_file(SRC_URL.format(reg))
        print(f"  reading {os.path.basename(path)}")
        rows = xlsx_rows(path)
        header = [str(h).strip() for h in next(rows)]
        ix = {h: i for i, h in enumerate(header)}
        i_tot = ix["Total Cancer Risk (per million)"]
        cols = header[i_tot + 1:]
        if src_cols is None:
            src_cols = cols
        assert cols == src_cols, (reg, cols)
        i_blk, i_pop = ix["Block"], ix["Population"]
        for r in rows:
            if r[i_blk] in (None, ""):
                continue
            b = block_id(r[i_blk])
            pop = num(r[i_pop])
            vals = [num(v) for v in r[i_tot + 1:i_tot + 1 + len(cols)]]
            tot = sum(vals)
            checks["blocks"] += 1
            pub = num(r[i_tot])
            # EPA rounds the published total to 1 significant figure
            if pub > 0:
                mag = 10 ** math.floor(math.log10(pub))
                checks["max_abs_diff_vs_rounded_total"] = max(
                    checks["max_abs_diff_vs_rounded_total"], abs(tot - pub) / mag)
            g = b[:11]
            t = tracts.setdefault(g, [0.0, 0.0, 0])
            t[0] += pop
            t[1] += pop * tot
            t[2] += 1
            if g in cd1_ids:
                d = detail.setdefault(g, defaultdict(float))
                for c, v in zip(cols, vals):
                    d[c] += pop * v
    out = {
        "tracts": {g: {"pop": p, "risk": (s / p if p else None), "n_blocks": n}
                   for g, (p, s, n) in tracts.items()},
        "cd1_detail": {g: {c: v / tracts[g][0] for c, v in d.items()}
                       for g, d in detail.items() if tracts[g][0]},
        "regions": REGIONS,
        "source_columns": src_cols,
        "checks": checks,
    }
    with open(cache, "w") as f:
        json.dump(out, f)
    return out


def pctile(values, x):
    """Percent of `values` below x, counting ties as half (0-100, integer)."""
    lo = sum(1 for v in values if v < x)
    eq = sum(1 for v in values if v == x)
    return round(100.0 * (lo + 0.5 * eq) / len(values))


# --------------------------------------------------------------------------
# CD1: cancer risk by pollutant
# --------------------------------------------------------------------------

POLLUTANT_NAMES = {
    "PAHPOM": "PAH/POM",
    "BENZOAPYRENE": "Benzo(a)pyrene",
    "CHROMIUM VI (HEXAVALENT)": "Chromium VI",
    "ARSENIC COMPOUNDS(INORGANIC INCLUDING ARSINE)": "Arsenic compounds",
    "1,3-BUTADIENE": "1,3-Butadiene",
    "1,4-DICHLOROBENZENE": "1,4-Dichlorobenzene",
    "1,4-DIOXANE": "1,4-Dioxane",
    "ETHYLENE DICHLORIDE (1,2-DICHLOROETHANE)": "Ethylene dichloride",
    "ETHYLENE DIBROMIDE (DIBROMOETHANE)": "Ethylene dibromide",
    "POLYCHLORINATED BIPHENYLS (AROCLORS)": "PCBs",
    "BIS(2-ETHYLHEXYL)PHTHALATE (DEHP)": "DEHP",
}


def pollutant_name(col):
    """EPA's upper-case column name -> readable name (sentence case)."""
    name = col.replace("Cancer Risk (per million)", "").strip()
    return POLLUTANT_NAMES.get(name, name[:1] + name[1:].lower())


def cd1_pollutants(cd1_ids):
    rows = xlsx_rows(cached_file(POLL_URL))
    header = [str(h).strip() for h in next(rows)]
    ix = {h: i for i, h in enumerate(header)}
    i_tot = ix["Total Cancer Risk (per million)"]
    cols = header[i_tot + 1:]
    i_blk, i_pop = ix["Block"], ix["Population"]
    acc = {}
    for r in rows:
        if r[i_blk] in (None, "") or not str(r[i_blk]).startswith("36047"):
            continue
        b = block_id(r[i_blk])
        g = b[:11]
        if g not in cd1_ids:
            continue
        pop = num(r[i_pop])
        a = acc.setdefault(g, [0.0, defaultdict(float)])
        a[0] += pop
        for c, v in zip(cols, r[i_tot + 1:i_tot + 1 + len(cols)]):
            a[1][c] += pop * num(v)
    return {g: {pollutant_name(c): v / p for c, v in d.items()}
            for g, (p, d) in acc.items() if p}


# --------------------------------------------------------------------------
# CD1: respiratory hazard index from EPA exposure concentrations and RfCs
# --------------------------------------------------------------------------

def resp_rfcs():
    """{website pollutant name: RfC in mg/m3} for respiratory-endpoint air toxics,
    from EPA's 2020AirToxScreenPollutants.xlsx."""
    z = zipfile.ZipFile(cached_file(SUPP_URL))
    name = [n for n in z.namelist() if n.endswith("2020AirToxScreenPollutants.xlsx")][0]
    path = os.path.join(CACHE_DIR, "airtoxscreen2020_pollutants.xlsx")
    if not os.path.exists(path):
        with open(path, "wb") as f:
            f.write(z.read(name))
    rows = xlsx_rows(path)
    header = [str(h).strip() if h is not None else "" for h in next(rows)]
    ix = {h: i for i, h in enumerate(header)}
    i_name = ix["AirToxScreen WEBSITE POLL NAME (comment)"]
    out, ures = {}, {}
    for r in rows:
        nm = str(r[i_name] or "").strip()
        if not nm:
            continue
        rfc = num(r[ix["RFC (mg/m3)"]])
        ure = num(r[ix["URE 1/(ug/m3)"]])
        if ure:
            ures.setdefault(nm, set()).add(ure)
        if str(r[ix["resp"]] or "").strip() and rfc > 0:
            out.setdefault(nm, set()).add(rfc)
    for nm, s in list(out.items()):
        assert len(s) == 1, (nm, s)
        out[nm] = s.pop()
    return out, {k: v.pop() for k, v in ures.items() if len(v) == 1}


def norm(s):
    return "".join(ch for ch in s.upper() if ch.isalnum())


def cd1_resp_hi(cd1_ids, rfcs):
    rows = xlsx_rows(cached_file(EXP_URL))
    header = [str(h).strip() if h is not None else "" for h in next(rows)]
    ix = {h: i for i, h in enumerate(header)}
    i_blk, i_pop = ix["Block"], ix["Population"]
    # match exposure-concentration columns to the respiratory pollutants
    by_norm = {norm(h): h for h in header}
    use = {}
    for nm, rfc in rfcs.items():
        h = by_norm.get(norm(nm))
        if h is None:
            continue
        use[h] = rfc
    missing = sorted(set(rfcs) - {k for k in rfcs if norm(k) in by_norm})
    acc = {}
    for r in rows:
        if r[i_blk] in (None, "") or not str(r[i_blk]).startswith("36047"):
            continue
        b = block_id(r[i_blk])
        g = b[:11]
        if g not in cd1_ids:
            continue
        pop = num(r[i_pop])
        a = acc.setdefault(g, [0.0, defaultdict(float)])
        a[0] += pop
        for h, rfc in use.items():
            # exposure concentrations are in ug/m3; RfCs in mg/m3
            a[1][h] += pop * num(r[ix[h]]) / (rfc * 1000.0)
    hq = {g: {h: v / p for h, v in d.items()} for g, (p, d) in acc.items() if p}
    return hq, sorted(use), missing


# --------------------------------------------------------------------------

def fmt_top(d, n=3):
    top = sorted(d.items(), key=lambda kv: -kv[1])[:n]
    return ", ".join(f"{k} {v:.1f}" for k, v in top)


def main():
    with open(os.path.join(DATA_DIR, "acs.geojson")) as f:
        tracts = json.load(f)["features"]
    ids = {t["properties"]["geoid"] for t in tracts}
    assert len(ids) == 54, len(ids)

    nat = national_tracts(ids)
    T = nat["tracts"]
    us = [v["risk"] for g, v in T.items() if v["pop"] > 0 and g[:2] not in NON_STATES]
    ny = [v["risk"] for g, v in T.items() if v["pop"] > 0 and g[:2] == "36"]
    nyc = [v["risk"] for g, v in T.items() if v["pop"] > 0 and g[:5] in NYC_COUNTIES]
    kings = [(v["pop"], v["risk"]) for g, v in T.items() if v["pop"] > 0 and g[:5] == "36047"]
    us_pw = (sum(v["pop"] * v["risk"] for g, v in T.items()
                 if v["pop"] > 0 and g[:2] not in NON_STATES)
             / sum(v["pop"] for g, v in T.items() if v["pop"] > 0 and g[:2] not in NON_STATES))
    kings_pw = sum(p * r for p, r in kings) / sum(p for p, _ in kings)
    print(f"  {nat['checks']['blocks']} populated blocks; "
          f"{len(us)} US tracts (50 states + DC), {len(ny)} NY, {len(nyc)} NYC")

    pol = cd1_pollutants(ids)
    rfcs, ures = resp_rfcs()
    hq, used, missing = cd1_resp_hi(ids, rfcs)
    print(f"  respiratory HI from {len(used)} pollutants; no exposure column for: {missing}")

    feats = []
    for t in tracts:
        p = t["properties"]
        g = p["geoid"]
        info = T.get(g)
        has = bool(info and info["pop"] > 0)
        props = {"geoid": g, "tract_label": p["tract_label"], "nta_name": p["nta_name"]}
        if has:
            risk = info["risk"]
            det = nat["cd1_detail"][g]
            grp = defaultdict(float)
            for c, v in det.items():
                grp[group_of(c)] += v
            hi = sum(hq[g].values())
            props.update({
                "cancer_risk_per_million": round(risk, 1),
                "cancer_risk_pctile_national": pctile(us, risk),
                "cancer_risk_pctile_state": pctile(ny, risk),
                "respiratory_hi": round(hi, 2),
                "top_cancer_pollutants": fmt_top(pol[g]),
                "top_cancer_sources": fmt_top(grp),
                # the same top-3 at EPA's detailed source-type level, leaving out
                # secondary formation and background (nearly uniform across CD1)
                "top_cancer_source_types": fmt_top({
                    source_label(c): v for c, v in det.items()
                    if group_of(c) not in ("Secondary formation", "Background")}),
                "cancer_risk_by_source_group": {k: round(grp[k], 2) for _, k in GROUPS},
                "top_respiratory_pollutants": ", ".join(
                    f"{k.title()} {v:.2f}" for k, v in
                    sorted(hq[g].items(), key=lambda kv: -kv[1])[:3]),
                "population_2020": int(round(info["pop"])),
                "n_blocks": info["n_blocks"],
                "no_population": False,
            })
        else:
            props.update({
                "cancer_risk_per_million": None, "cancer_risk_pctile_national": None,
                "cancer_risk_pctile_state": None, "respiratory_hi": None,
                "top_cancer_pollutants": None, "top_cancer_sources": None,
                "top_cancer_source_types": None,
                "cancer_risk_by_source_group": None, "top_respiratory_pollutants": None,
                "population_2020": int(round(info["pop"])) if info else 0,
                "n_blocks": info["n_blocks"] if info else 0,
                "no_population": True,
            })
        props["assessment_year"] = YEAR
        props["tract_vintage"] = 2020
        feats.append({"type": "Feature", "properties": props, "geometry": t["geometry"]})

    vals = [f["properties"] for f in feats if not f["properties"]["no_population"]]
    cr = sorted(v["cancer_risk_per_million"] for v in vals)
    hi = sorted(v["respiratory_hi"] for v in vals)
    print(f"  cancer risk {cr[0]}-{cr[-1]} per million; respiratory HI {hi[0]}-{hi[-1]}")
    print(f"  tract medians: US {statistics.median(us):.1f}, NY {statistics.median(ny):.1f}, "
          f"NYC {statistics.median(nyc):.1f}; pop-weighted US {us_pw:.1f}, Kings {kings_pw:.1f}")

    meta = {
        "source": "U.S. EPA AirToxScreen 2020 (block-level results, gaftp.epa.gov/"
                  "rtrmodeling_public/AirToxScreen/2020), aggregated to 2020 census tracts; "
                  "geometry from data/acs.geojson",
        "assessment_year": YEAR,
        "notes": "Screening-level modeled estimates. cancer_risk_per_million = population-"
                 "weighted mean over the tract's populated 2020 blocks of EPA's total "
                 "lifetime cancer risk (sum of all source types incl. secondary formation and "
                 "background; diesel PM not included). respiratory_hi = population-weighted "
                 "sum of EPA exposure concentration / RfC over EPA's respiratory-endpoint air "
                 "toxics (computed here; EPA has not published 2020 HI files). Percentiles "
                 "rank the tract among populated tracts (US = 50 states + DC).",
        # round-number 5-class breaks chosen for the 2020 CD1 values (lower edge inclusive)
        "suggested_breaks": {
            "cancer_risk_per_million": [33, 35, 36, 37, 39, 41],
            "respiratory_hi": [0.40, 0.45, 0.47, 0.49, 0.51, 0.55],
        },
        "context": {
            "us_tract_median": round(statistics.median(us), 1),
            "ny_tract_median": round(statistics.median(ny), 1),
            "nyc_tract_median": round(statistics.median(nyc), 1),
            "us_population_weighted_mean": round(us_pw, 1),
            "kings_population_weighted_mean": round(kings_pw, 1),
        },
    }
    write_fc("air_toxics_risk.geojson", feats, meta)


if __name__ == "__main__":
    main()
